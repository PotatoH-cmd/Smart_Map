"""对话主链路：/chat · /chat/stream · Run 生命周期 · 会话管理（自 main.py 机械搬移，行为不变）。"""
from pydantic import BaseModel
from typing import List, Dict, Any, Optional
import httpx
import os
import base64
from fastapi import FastAPI, APIRouter, HTTPException, Header, Request, UploadFile, File, Form, Query
from prompts import build_view_system_message, build_elevation_injection
import json
import traceback
from fastapi.responses import StreamingResponse, JSONResponse, Response, FileResponse, HTMLResponse
from agents.context_manager import load_history_from_db, COMPRESS_THRESHOLD_TURNS
import asyncio
from agents.run_engine import is_confirm_message, parse_supplied_from_message
import uuid
from agents.run_store import get_run_store
import contextlib
import re
import logging
import main as _main  # 仅访问 lifespan 装配的单例：_main.task_executor / _main.bot
from core.db import DB_PATH, get_db, now_iso

# ── OCR / 上传 / SSE 常量（自 main.py 收敛到 chat 域）──
OCR_MODEL = 'qwen-vl-ocr-2025-11-20'
OCR_BASE_URL = 'https://dashscope.aliyuncs.com/compatible-mode/v1'
OCR_API_KEY = os.environ.get('DASHSCOPE_API_KEY', '')
UPLOAD_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "static")  # 图片本地兜底目录
RUN_SSE_IDLE_TIMEOUT = 120  # run SSE 订阅空闲超时：超过该时长无事件则检查 run 状态（防哨兵丢失导致挂起）


def _is_explicit_map_request(user_content: str) -> bool:
    if not user_content:
        return False
    map_terms = [
        "地图", "上图", "加载", "图层", "定位", "跳转", "飞到", "显示到地图",
        "加载到地图", "标记", "标注", "打点", "落点", "经纬度", "卫星图", "底图",
        "切换", "卫星", "清除",
        # qgis_mcp_tool 产出结果也需要地图展示
        "中心点", "缓冲区", "buffer", "裁剪", "clip",
        # 空间分析类结果也需要地图展示
        "距离", "连线", "最近", "最短", "多远",
    ]
    return any(term in user_content for term in map_terms)


def _is_explicit_marker_request(user_content: str) -> bool:
    if not user_content:
        return False
    marker_terms = ["标记", "标注", "打点", "落点", "经纬度", "坐标点"]
    return any(term in user_content for term in marker_terms)


def _stable_json_key(payload: Any) -> str:
    try:
        return json.dumps(payload, ensure_ascii=False, sort_keys=True)
    except Exception:
        return str(payload)


logger = logging.getLogger(__name__)

router = APIRouter()


class ChatMessage(BaseModel):
    role: str
    content: str
    images: Optional[List[str]] = None  # 图片 URL 列表（用于多模态消息）
class ChatRequest(BaseModel):
    messages: List[ChatMessage]
    active_view: str = 'map'  # 'map'(2D) | 'cesium'(3D) | 'kb'
    session_id: Optional[str] = None  # 可选会话 ID，用于 MemorySaver thread_id 隔离
class ChatResponse(BaseModel):
    response: str
    messages: List[Dict[str, Any]]
    map_commands: List[Dict[str, Any]] = []
    cesium_commands: List[Dict[str, Any]] = []
    charts: List[Dict[str, Any]] = []
    report_url: Optional[str] = None
    intent_info: Optional[Dict[str, Any]] = None
def _extract_text_content(content) -> str:
    """从消息内容中提取纯文本。若为多模态数组，取第一个 text 部分。"""
    if isinstance(content, list):
        for part in content:
            if isinstance(part, dict) and part.get("type") == "text":
                return part.get("text", "")
        return ""
    return content or ""
async def _ocr_image(image_url: str) -> str:
    """使用 qwen-vl-ocr 专用模型从图片中提取文字内容。
    
    image_url 可以是：
    - 本地路径（如 /static/uploads/xxx.png），会自动转为 base64 data URL
    - 完整 HTTP(S) URL（DashScope 需要能访问）
    """
    # 将本地路径转为 base64 data URL（DashScope 无法访问内网地址）
    if image_url.startswith('/'):
        local_path = os.path.join(os.path.dirname(__file__), image_url.lstrip('/'))
        if not os.path.isfile(local_path):
            # 尝试相对于 static 目录
            local_path = os.path.join(UPLOAD_DIR, image_url.lstrip('/'))
        if os.path.isfile(local_path):
            with open(local_path, 'rb') as f:
                img_data = f.read()
            # 探测 MIME 类型
            ext = os.path.splitext(local_path)[1].lower()
            mime_map = {'.png': 'image/png', '.jpg': 'image/jpeg', '.jpeg': 'image/jpeg',
                        '.gif': 'image/gif', '.webp': 'image/webp', '.bmp': 'image/bmp'}
            mime = mime_map.get(ext, 'image/png')
            image_url = f"data:{mime};base64,{base64.b64encode(img_data).decode()}"
            logger.info(f"[OCR] 本地图片转 base64: {local_path} ({len(img_data)} bytes)")
        else:
            logger.error(f"[OCR] 图片文件不存在: {local_path}")
            return ""
    
    async with httpx.AsyncClient(timeout=120.0) as client:
        response = await client.post(
            f"{OCR_BASE_URL}/chat/completions",
            headers={
                "Authorization": f"Bearer {OCR_API_KEY}",
                "Content-Type": "application/json",
            },
            json={
                "model": OCR_MODEL,
                "messages": [
                    {
                        "role": "system",
                        "content": "你是一个精确的文字提取工具。只输出图片中的纯文字内容，不要添加任何坐标框、边框标记、位置标注或额外说明。保持原始格式、精度和顺序。"
                    },
                    {
                        "role": "user",
                        "content": [
                            {
                                "type": "text",
                                "text": "提取图片中所有文字，只输出纯文本，不要任何坐标标注或边框信息。"
                            },
                            {
                                "type": "image_url",
                                "image_url": {"url": image_url},
                            }
                        ]
                    }
                ],
                "max_tokens": 4000,
            }
        )
        if response.status_code != 200:
            logger.error(f"[OCR] API 请求失败 HTTP {response.status_code}: {response.text[:300]}")
            return ""
        result = response.json()
        text = result["choices"][0]["message"]["content"]
        logger.info(f"[OCR] 提取成功 ({len(text)} 字符): {text[:200]}...")
        return text
async def _preprocess_excel(file_url: str) -> str:
    """使用 openpyxl 解析 Excel 文件，将表格数据格式化为文本。"""
    local_path = os.path.join(os.path.dirname(__file__), file_url.lstrip('/'))
    if not os.path.isfile(local_path):
        # 尝试相对于 static 目录
        local_path = os.path.join(UPLOAD_DIR, file_url.lstrip('/'))
    if not os.path.isfile(local_path):
        logger.error(f"[Excel] 文件不存在: {file_url} -> {local_path}")
        return "[Excel 文件读取失败]"
    try:
        import openpyxl
        wb = openpyxl.load_workbook(local_path, data_only=True)
        parts = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            if ws.max_row == 0:
                continue
            rows = []
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 200), values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                rows.append(" | ".join(cells))
            if rows:
                parts.append(f"【Excel 文件内容（{sheet_name}）】\n" + "\n".join(rows))
        wb.close()
        if parts:
            result = "\n\n".join(parts)
            logger.info(f"[Excel] 解析成功: {file_url} ({ws.max_row} 行)")
            return result
        return "[Excel 文件为空]"
    except Exception as e:
        logger.error(f"[Excel] 解析失败: {e}")
        return f"[Excel 文件解析失败: {e}]"
async def _preprocess_message_images(messages: list) -> list:
    """对用户消息中的附件进行预处理。
    
    - 图片：用 qwen-vl-ocr 提取文字
    - Excel：用 openpyxl 解析表格
    预处理结果注入到消息 content 中，并移除 images 字段。
    """
    for msg in messages:
        attachments = msg.get('images')
        if msg.get('role') != 'user' or not attachments or not isinstance(attachments, list):
            continue

        text_content = msg.get('content', '')
        all_results = []
        for file_url in attachments:
            is_excel = file_url.lower().endswith(('.xlsx', '.xls'))
            is_image = file_url.lower().endswith(('.png', '.jpg', '.jpeg', '.gif', '.webp', '.bmp'))
            
            if is_excel:
                logger.info(f"[Excel] 开始处理: {file_url}")
                try:
                    result = await _preprocess_excel(file_url)
                    if result.strip():
                        all_results.append(result)
                except Exception as e:
                    logger.error(f"[Excel] 处理异常: {e}")
                    all_results.append(f"[Excel 读取失败: {e}]")
            elif is_image:
                logger.info(f"[OCR] 开始处理图片: {file_url}")
                try:
                    ocr_text = await _ocr_image(file_url)
                    if ocr_text.strip():
                        all_results.append(ocr_text)
                except Exception as e:
                    logger.error(f"[OCR] 图片处理异常: {e}")
                    all_results.append(f"[图片 OCR 失败: {e}]")
            else:
                logger.warning(f"[预处理] 跳过不支持的文件: {file_url}")

        if all_results:
            prefix = "【以下是从上传文件中自动提取的内容】\n"
            msg['content'] = text_content + "\n\n" + prefix + "\n---\n".join(all_results)
            logger.info(f"[预处理] 已将 {len(all_results)} 个文件内容注入到用户消息")
        
        # 移除 images 字段——主模型 qwen-flash 不支持，且文件内容已注入
        msg.pop('images', None)
    
    return messages
@router.post("/chat")
async def chat(request: ChatRequest, x_session_id: Optional[str] = Header(default=None, alias="X-Session-ID")):
    try:
        if not request.messages:
            raise HTTPException(status_code=400, detail="消息不能为空")

        # 优先使用请求头 X-Session-ID，其次 body.session_id，最后回退 "default"
        thread_id = x_session_id or request.session_id or "default"
        
        valid_roles = {'user', 'assistant', 'system', 'function'}
        for i, msg in enumerate(request.messages):
            if msg.role not in valid_roles:
                raise HTTPException(status_code=400, detail=f"无效角色: {msg.role}")
        
        messages = [msg.model_dump() for msg in request.messages]
        # 记录每条消息的角色，用于调试 400 错误
        roles = [m.get('role') for m in messages]
        logger.info(f"Processing messages with roles: {roles}")
        
        # 修复：确保消息列表以 user 开头（跳过 system）
        # 如果第一条不是 system 且第一条不是 user，或者第一条是 system 且第二条不是 user
        if messages:
            first_non_system_idx = 0
            if messages[0].get('role') == 'system':
                first_non_system_idx = 1
            
            if len(messages) > first_non_system_idx:
                if messages[first_non_system_idx].get('role') != 'user':
                    logger.warning(f"First non-system message is {messages[first_non_system_idx].get('role')}, not 'user'. Attempting to fix...")
                    # 找到第一个 user 消息并删除它之前的所有非 system 消息
                    first_user_idx = -1
                    for i in range(first_non_system_idx, len(messages)):
                        if messages[i].get('role') == 'user':
                            first_user_idx = i
                            break
                    
                    if first_user_idx != -1:
                        messages = messages[:first_non_system_idx] + messages[first_user_idx:]
                        logger.info(f"Fixed message sequence. New roles: {[m.get('role') for m in messages]}")
                    else:
                        logger.error("No user message found in history!")

        # 强力指令：彻底纠正实测高程与超深逻辑，严禁误导
        if messages and messages[-1]['role'] == 'user':
            content = messages[-1]['content']
            if any(k in content for k in ["超深", "高程", "深度", "采深", "开采深度"]):
                # SSoT：高程业务规则唯一权威在 prompts.py::build_elevation_injection
                messages[-1]['content'] += build_elevation_injection()
                logger.info(f"Injected elevation business rules.")
            # 地图数据加载意图守卫：根据 active_view 注入 2D/3D 视图提示词（SSoT：prompts.py）
            map_intent_keywords = ["加载", "上图", "矢量", "图层", "地图", "位置", "显示到地图", "加载到地图", "跳转", "定位", "标记", "标点", "经纬度", "切换", "卫星", "底图", "清除"]
            if any(k in content for k in map_intent_keywords):
                messages.insert(0, {
                    'role': 'system',
                    'content': build_view_system_message(request.active_view),
                })
                logger.info(f"{'3D' if request.active_view == 'cesium' else '2D'} view detected: view prompt injected.")

        use_intent_agent = os.environ.get("USE_INTENT_AGENT", "true").lower() == "true"

        if use_intent_agent and _main.task_executor is not None:
            result = await _main.task_executor.execute(
                user_message=messages[-1]['content'],
                chat_history=messages[:-1],
                thread_id=thread_id,
            )

            intent_info = None
            if result.get("intent_result"):
                ir = result["intent_result"]
                intent_info = {
                    "primary_intent": ir.primary_intent.value if hasattr(ir.primary_intent, 'value') else str(ir.primary_intent),
                    "confidence": ir.confidence,
                    "task_context": ir.task_context,
                    "entities": ir.entities,
                    "execution_plan": [
                        {
                            "step_id": s.step_id,
                            "action": s.action,
                            "tool": s.tool,
                            "reasoning": s.reasoning
                        }
                        for s in ir.execution_plan
                    ] if ir.execution_plan else []
                }

            # 持久化 intent agent 结果
            _persist_chat(thread_id, messages[-1]['content'], result.get("response", ""))
            _schedule_fact_extraction(thread_id, messages[-1]['content'], result.get("response", ""))

            optimized_map_commands = _optimize_map_commands(messages[-1]['content'], result.get("map_commands", []))
            optimized_charts = _optimize_charts(result.get("charts", []))

            return ChatResponse(
                response=result.get("response", "命令已执行。"),
                messages=result.get("messages", []),
                map_commands=optimized_map_commands,
                cesium_commands=result.get("cesium_commands", []),
                charts=optimized_charts,
                report_url=result.get("report_url"),
                intent_info=intent_info
            )

        response_messages = []
        iteration = 0
        runner = _main.bot
        if messages and messages[-1]['role'] == 'user':
            content = messages[-1]['content']
            weather_intent_keywords = ["天气", "气温", "几度", "空气质量", "AQI", "雾霾", "降雨", "下雨", "预报", "带伞", "风力", "风速", "湿度", "紫外线"]
            if any(k in content for k in weather_intent_keywords):
                runner = _main.bot  # 主 _main.bot 已有 weather_tool
                logger.info("Weather intent: using main _main.bot.")
            else:
                map_intent_keywords = ["加载", "上图", "矢量", "图层", "地图", "位置", "显示到地图", "加载到地图", "跳转", "定位", "标记", "标点", "经纬度"]
                if any(k in content for k in map_intent_keywords):
                    # 始终使用主 _main.bot（已注册所有工具），通过系统消息指导工具选择
                    # 不使用 build_assistant_with_tools 的受限 Assistant，避免工具注册问题
                    runner = _main.bot
                    if request.active_view == 'cesium':
                        logger.info("3D view: using main _main.bot with cesium_tool system hint.")
                    else:
                        logger.info("2D view: using main _main.bot with map_tool system hint.")
        for response in runner.run(messages=messages):
            iteration += 1
            response_messages = response
            # 记录每一次迭代，看是否卡在某个工具调用上
            last_msg = response_messages[-1] if response_messages else {}
            role = last_msg.get('role')
            name = last_msg.get('name', 'N/A')
            content = last_msg.get('content', '')
            
            # 记录关键信息
            if role == 'assistant' and 'call' in str(last_msg):
                logger.info(f"Bot iteration {iteration}: Assistant is calling tool: {last_msg}")
            elif role == 'function':
                logger.info(f"Bot iteration {iteration}: Function {name} returned: {str(content)[:200]}...")
            else:
                logger.info(f"Bot iteration {iteration}: role={role}, content={str(content)[:100]}...")
        
        # 获取最后一条助手回复作为 response 字段（兼容旧前端逻辑）
        response_text = ""
        formatted_answer = ""
        for msg in reversed(response_messages):
            if msg.get('role') == 'assistant' and msg.get('content'):
                response_text = clean_response_content(msg['content'])
                break
        # 如果工具返回了标准化答案（如平均超深深度），优先使用
        if not response_text:
            for msg in reversed(response_messages):
                if msg.get('role') == 'function' and msg.get('name') == 'data_visualizer_tool':
                    try:
                        func_res = json.loads(msg.get('content', '{}'))
                        fa = func_res.get('formatted_answer')
                        if fa:
                            formatted_answer = fa
                            break
                    except:
                        pass
        if formatted_answer:
            response_text = formatted_answer
        
        if not response_text:
            response_text = "命令已执行。"

        # 提取地图命令
        map_commands = []
        cesium_commands = []
        charts = []
        for msg in response_messages:
            if msg.get('role') == 'function' and msg.get('name') in ['map_tool', 'location_search']:
                try:
                    func_res = json.loads(msg.get('content', '{}'))
                    if func_res.get('map_command'):
                        map_commands.append(func_res['map_command'])
                except:
                    pass
            # 提取 Cesium 3D 命令
            if msg.get('role') == 'function' and msg.get('name') == 'cesium_tool':
                try:
                    func_res = json.loads(msg.get('content', '{}'))
                    if func_res.get('cesium_command'):
                        cesium_commands.append(func_res['cesium_command'])
                except:
                    pass
            # 提取数据可视化图表
            if msg.get('role') == 'function' and msg.get('name') == 'data_visualizer_tool':
                try:
                    vis = json.loads(msg.get('content', '{}'))
                    if isinstance(vis, dict) and vis.get('success'):
                        charts.append({
                            'chart_type': vis.get('chart_type'),
                            'config': vis.get('config'),
                            'summary': vis.get('content')
                        })
                except Exception as e:
                    logger.warning(f"Failed to parse visualization content: {e}")
        
        map_commands = _optimize_map_commands(messages[-1]['content'], map_commands)
        charts = _optimize_charts(charts)

        logger.info(f"Chat response: {response_text[:100]}... Map commands: {len(map_commands)}, Cesium commands: {len(cesium_commands)}, Charts: {len(charts)}")
        if map_commands:
            logger.info(f"First map command: {map_commands[0]}")
        if cesium_commands:
            logger.info(f"First cesium command: {cesium_commands[0]}")

        # 持久化消息到 SQLite（仅当 thread_id 是有效 UUID 会话时）
        _persist_chat(thread_id, messages[-1]['content'], response_text)
        _schedule_fact_extraction(thread_id, messages[-1]['content'], response_text)

        return ChatResponse(
            response=response_text,
            messages=response_messages, # 返回完整的对话历史，包括工具执行结果
            map_commands=map_commands,
            cesium_commands=cesium_commands,
            charts=charts
        )
    
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Chat error: {e}\n{traceback.format_exc()}")
        raise HTTPException(status_code=500, detail="内部服务器错误")
@router.post("/chat/stream")
async def chat_stream(request: ChatRequest, x_session_id: Optional[str] = Header(default=None, alias="X-Session-ID"), x_pending_run_id: Optional[str] = Header(default=None, alias="X-Pending-Run-ID")):
    def sse_payload(payload: Dict[str, Any]) -> str:
        return f"data: {json.dumps(payload, ensure_ascii=False)}\n\n"

    async def event_generator():
        try:
            # 先发送 SSE 预热包，尽量让浏览器和代理尽快进入流式模式
            yield ": stream-open\n\n"
            yield f": {' ' * 2048}\n\n"
            yield "retry: 1500\n\n"

            if not request.messages:
                raise HTTPException(status_code=400, detail="消息不能为空")

            thread_id = x_session_id or request.session_id or "default"
            valid_roles = {'user', 'assistant', 'system', 'function'}
            for msg in request.messages:
                if msg.role not in valid_roles:
                    raise HTTPException(status_code=400, detail=f"无效角色: {msg.role}")

            messages = [msg.model_dump() for msg in request.messages]
            roles = [m.get('role') for m in messages]
            logger.info(f"[stream] Processing messages with roles: {roles}")

            # 阶段4 契约改造：前端只传 session_id + 本轮消息时，历史由后端从 DB 读取并裁剪
            # （≤COMPRESS_THRESHOLD_TURNS 轮），避免每轮全量回传造成上下文污染。
            if (len(messages) == 1 and messages[0].get('role') == 'user'
                    and thread_id and thread_id != "default"):
                db_history = load_history_from_db(DB_PATH, thread_id)
                if db_history:
                    # DB 历史与本轮消息不重叠（_persist_chat 在 final 时才写入）
                    messages = db_history + messages
                    logger.info(f"[stream] history loaded from db: {len(db_history)} msgs (session={thread_id})")

            # 阶段4：超长会话滚动摘要压缩（fire-and-forget，失败静默降级为截断）
            if _main.task_executor is not None and getattr(_main.task_executor, "context_manager", None) is not None:
                try:
                    asyncio.create_task(_main.task_executor.context_manager.compress_history(
                        thread_id, messages, getattr(_main.task_executor, "_llm", None),
                    ))
                except Exception as e:
                    logger.warning(f"[stream] compress_history task failed: {e}")

            # 图片 OCR 预处理：用 qwen-vl-ocr 提取文字后注入到用户消息
            # 主模型 qwen-flash 不支持图片，所以文字提取后再移除 images 字段
            messages = await _preprocess_message_images(messages)

            if messages:
                first_non_system_idx = 0
                if messages[0].get('role') == 'system':
                    first_non_system_idx = 1

                if len(messages) > first_non_system_idx and messages[first_non_system_idx].get('role') != 'user':
                    first_user_idx = -1
                    for i in range(first_non_system_idx, len(messages)):
                        if messages[i].get('role') == 'user':
                            first_user_idx = i
                            break
                    if first_user_idx != -1:
                        messages = messages[:first_non_system_idx] + messages[first_user_idx:]

            raw_user_text = None
            if messages and messages[-1]['role'] == 'user':
                content = _extract_text_content(messages[-1]['content'])
                # 保存注入前的原始用户文本：高程业务规则注入会向消息追加含快速路由
                # 关键词的文本（如"红线标准"），快速路由/意图分析必须基于纯用户输入，
                # 否则"XX可采区的高程是多少"会被注入文本误导为 spatial_reference。
                raw_user_text = content
                if any(k in content for k in ["超深", "高程", "深度", "采深", "开采深度"]):
                    # 多模态消息需要特殊处理：将业务规则追加到文本部分
                    # SSoT：高程业务规则唯一权威在 prompts.py::build_elevation_injection
                    extra_rules = build_elevation_injection()
                    if isinstance(messages[-1]['content'], list):
                        # 多模态格式：追加到第一个 text 部分
                        for part in messages[-1]['content']:
                            if isinstance(part, dict) and part.get("type") == "text":
                                part["text"] = part.get("text", "") + extra_rules
                                break
                    else:
                        messages[-1]['content'] += extra_rules
                map_intent_keywords = ["加载", "上图", "矢量", "图层", "地图", "位置", "显示到地图", "加载到地图", "跳转", "定位", "标记", "标点", "经纬度", "切换", "卫星", "底图", "清除"]
                if any(k in content for k in map_intent_keywords):
                    # 视图提示词 SSoT：prompts.py::build_view_system_message（阶段4/6）
                    messages.insert(0, {
                        'role': 'system',
                        'content': build_view_system_message(request.active_view),
                    })

            yield sse_payload({
                "type": "status",
                "stage": "queued",
                "message": "请求已提交，后端开始处理。",
            })

            use_intent_agent = os.environ.get("USE_INTENT_AGENT", "true").lower() == "true"
            run_engine_on = os.environ.get("RUN_ENGINE", "on").lower() in ("1", "on", "true")

            # ── 阶段1新路径：RunEngine（灰度开关 RUN_ENGINE，默认 on）──
            # run 在独立 asyncio task 中执行，SSE 只是事件订阅者；断线不影响执行
            if run_engine_on and _main.task_executor is not None and _main.task_executor.run_engine is not None:
                user_text = raw_user_text or _extract_text_content(messages[-1].get('content', ''))
                engine = _main.task_executor.run_engine
                store = engine.store
                bus = engine.bus

                # pending resume 判定（规则化：确认 / 补参，不调 LLM）
                run_id = None
                task = None
                pending = {}
                pending_run = None
                resume_source_text = None  # resume 场景的原始请求文本（用于地图命令优化判定）
                if x_pending_run_id:
                    pending_run = store.get_run(x_pending_run_id)
                    if pending_run and pending_run.get("status") in ("awaiting_confirmation", "awaiting_input"):
                        try:
                            pending = json.loads(pending_run.get("pending_json") or "{}")
                        except Exception:
                            pending = {}
                    else:
                        pending_run = None
                if pending_run is None:
                    pending_run = store.get_pending_by_session(thread_id)
                    pending = (pending_run or {}).get("pending") or {}

                if pending_run:
                    rid = pending_run["run_id"]
                    ptype = (pending or {}).get("pending_type", "")
                    if ptype == "confirm" and is_confirm_message(user_text):
                        task = await engine.resume(rid, confirm=True)
                        run_id = rid
                        logger.info(f"[stream] resume confirm run={rid}")
                    elif ptype == "input":
                        supplied = parse_supplied_from_message(pending, user_text)
                        if supplied:
                            task = await engine.resume(rid, user_supplied=supplied)
                            run_id = rid
                            logger.info(f"[stream] resume input run={rid} supplied={list(supplied.keys())}")
                    if task is not None:
                        # 补参/确认消息本身不含地图关键词（如"ceshen"），但原始请求
                        # 已判定为地图意图；用 checkpoint 中的原始 user_message 做
                        # 地图命令优化判定，避免 resume 场景误抑制 map_commands。
                        ckpt = store.load_checkpoint(rid)
                        resume_source_text = (ckpt or {}).get("user_message") or None
                    else:
                        logger.info(f"[stream] pending run={rid} 未匹配 resume 规则，另起新 run")

                if task is None:
                    run_id = str(uuid.uuid4())
                    task = await engine.start(
                        run_id=run_id, session_id=thread_id, user_message=user_text,
                        chat_history=messages[:-1],
                        view_hint=request.active_view,
                    )

                yield sse_payload({
                    "type": "status", "stage": "run_started",
                    "message": "请求已提交，run 已启动。",
                    "run_id": run_id,
                })

                # 订阅事件流并转 SSE。known_seq 为订阅前的历史边界：
                # 回放中的旧 final（如上一生命周期 pending 终结事件，resume 场景）
                # 不应终止本次 SSE，跳过它们等待真正的最终结果。
                known_seq = await bus.current_seq(run_id)
                queue = await bus.subscribe(run_id)
                try:
                    while True:
                        try:
                            event = await asyncio.wait_for(queue.get(), timeout=RUN_SSE_IDLE_TIMEOUT)
                        except asyncio.TimeoutError:
                            run = store.get_run(run_id)
                            if run and run["status"] in ("completed", "failed", "cancelled"):
                                break
                            continue
                        if event is None:
                            break
                        if event.get("type") == "final":
                            result = event.get("result", {})
                            if result.get("pending_type") and event.get("seq", 0) <= known_seq:
                                # 回放的旧 pending 终结事件：跳过，等新生命周期的事件
                                continue
                            result["run_id"] = run_id
                            result["map_commands"] = _optimize_map_commands(
                                resume_source_text or user_text, result.get("map_commands", [])
                            )
                            result["charts"] = _optimize_charts(result.get("charts", []))
                            _persist_chat(thread_id, user_text, result.get("response", ""))
                            _schedule_fact_extraction(thread_id, user_text, result.get("response", ""))
                            yield sse_payload({"type": "final", "result": result})
                            break
                        yield sse_payload(event)
                finally:
                    bus.unsubscribe(run_id, queue)

                yield "data: [DONE]\n\n"
                return

            if use_intent_agent and _main.task_executor is not None:
                # 旧路径：execute_stream 手动编排（RUN_ENGINE=off 时保留）
                # 提取纯文本用于 _main.task_executor（它目前只支持字符串 user_message）
                user_text = _extract_text_content(messages[-1].get('content', ''))
                async for event in _main.task_executor.execute_stream(
                    user_message=user_text,
                    chat_history=messages[:-1],
                    thread_id=thread_id,
                ):
                    if event.get("type") == "final":
                        result = event.get("result", {})
                        result["map_commands"] = _optimize_map_commands(user_text, result.get("map_commands", []))
                        result["charts"] = _optimize_charts(result.get("charts", []))
                        _persist_chat(thread_id, user_text, result.get("response", ""))
                        _schedule_fact_extraction(thread_id, user_text, result.get("response", ""))
                        yield sse_payload({"type": "final", "result": result})
                    else:
                        yield sse_payload(event)

                yield "data: [DONE]\n\n"
                return

            response_messages = []
            iteration = 0
            runner = _main.bot
            if messages and messages[-1]['role'] == 'user':
                content = messages[-1]['content']
                weather_intent_keywords = ["天气", "气温", "几度", "空气质量", "AQI", "雾霾", "降雨", "下雨", "预报", "带伞", "风力", "风速", "湿度", "紫外线"]
                if any(k in content for k in weather_intent_keywords):
                    runner = _main.bot
                else:
                    map_intent_keywords = ["加载", "上图", "矢量", "图层", "地图", "位置", "显示到地图", "加载到地图", "跳转", "定位", "标记", "标点", "经纬度"]
                    if any(k in content for k in map_intent_keywords):
                        runner = _main.bot

            yield sse_payload({
                "type": "status",
                "stage": "model",
                "message": "模型已开始执行，正在逐步调用工具。",
            })

            for response in runner.run(messages=messages):
                iteration += 1
                response_messages = response
                last_msg = response_messages[-1] if response_messages else {}
                role = last_msg.get('role')
                name = last_msg.get('name', 'N/A')
                content = last_msg.get('content', '')

                if role == 'assistant' and 'call' in str(last_msg):
                    yield sse_payload({
                        "type": "tool_start",
                        "stage": "tool_start",
                        "tool_name": name,
                        "message": f"正在调用工具处理请求（第 {iteration} 轮）。",
                    })
                elif role == 'function':
                    try:
                        parsed = json.loads(content) if isinstance(content, str) else content
                    except Exception:
                        parsed = {"content": str(content)}
                    summary = parsed.get('content') or parsed.get('message') or parsed.get('error') or f"{name} 已返回结果"
                    yield sse_payload({
                        "type": "tool_result",
                        "stage": "tool_result",
                        "tool_name": name,
                        "message": f"{name}：{str(summary)[:100]}",
                    })
                elif role == 'assistant' and content:
                    preview = clean_response_content(str(content))[:80]
                    if preview:
                        yield sse_payload({
                            "type": "status",
                            "stage": "reasoning",
                            "message": f"正在整理回复：{preview}",
                        })

            response_text = ""
            formatted_answer = ""
            for msg in reversed(response_messages):
                if msg.get('role') == 'assistant' and msg.get('content'):
                    response_text = clean_response_content(msg['content'])
                    break
            if not response_text:
                for msg in reversed(response_messages):
                    if msg.get('role') == 'function' and msg.get('name') == 'data_visualizer_tool':
                        try:
                            func_res = json.loads(msg.get('content', '{}'))
                            fa = func_res.get('formatted_answer')
                            if fa:
                                formatted_answer = fa
                                break
                        except Exception:
                            pass
            if formatted_answer:
                response_text = formatted_answer
            if not response_text:
                response_text = "命令已执行。"

            map_commands = []
            cesium_commands = []
            charts = []
            for msg in response_messages:
                if msg.get('role') == 'function' and msg.get('name') in ['map_tool', 'location_search']:
                    try:
                        func_res = json.loads(msg.get('content', '{}'))
                        if func_res.get('map_command'):
                            map_commands.append(func_res['map_command'])
                    except Exception:
                        pass
                if msg.get('role') == 'function' and msg.get('name') == 'cesium_tool':
                    try:
                        func_res = json.loads(msg.get('content', '{}'))
                        if func_res.get('cesium_command'):
                            cesium_commands.append(func_res['cesium_command'])
                    except Exception:
                        pass
                if msg.get('role') == 'function' and msg.get('name') == 'data_visualizer_tool':
                    try:
                        vis = json.loads(msg.get('content', '{}'))
                        if isinstance(vis, dict) and vis.get('success'):
                            charts.append({
                                'chart_type': vis.get('chart_type'),
                                'config': vis.get('config'),
                                'summary': vis.get('content')
                            })
                    except Exception as e:
                        logger.warning(f"[stream] Failed to parse visualization content: {e}")

            map_commands = _optimize_map_commands(messages[-1]['content'], map_commands)
            charts = _optimize_charts(charts)
            _persist_chat(thread_id, messages[-1]['content'], response_text)
            _schedule_fact_extraction(thread_id, messages[-1]['content'], response_text)

            yield sse_payload({
                "type": "final",
                "result": {
                    "response": response_text,
                    "messages": response_messages,
                    "map_commands": map_commands,
                    "cesium_commands": cesium_commands,
                    "charts": charts,
                    "intent_info": None,
                }
            })
            yield "data: [DONE]\n\n"
        except HTTPException as e:
            message = e.detail if isinstance(e.detail, str) else "请求处理失败"
            # 失败轮次也落库，避免 DB 历史断档
            if messages:
                _persist_chat(thread_id, messages[-1]['content'], message)
            yield sse_payload({"type": "error", "stage": "error", "message": message})
            yield sse_payload({
                "type": "final",
                "result": {
                    "response": message,
                    "messages": [],
                    "map_commands": [],
                    "cesium_commands": [],
                    "charts": [],
                    "intent_info": None,
                }
            })
            yield "data: [DONE]\n\n"
        except Exception as e:
            logger.error(f"Chat stream error: {e}\n{traceback.format_exc()}")
            message = f"内部服务器错误：{str(e)}"
            # 失败轮次也落库，避免 DB 历史断档
            if messages:
                _persist_chat(thread_id, messages[-1]['content'], message)
            yield sse_payload({"type": "error", "stage": "error", "message": message})
            yield sse_payload({
                "type": "final",
                "result": {
                    "response": message,
                    "messages": [],
                    "map_commands": [],
                    "cesium_commands": [],
                    "charts": [],
                    "intent_info": None,
                }
            })
            yield "data: [DONE]\n\n"

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache, no-transform",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
            "Content-Type": "text/event-stream; charset=utf-8",
            "Content-Encoding": "identity",
        },
    )
@router.get("/api/run/{run_id}")
async def get_run_status(run_id: str):
    """查询 run 状态与 pending 载荷（断线重连 / 前端轮询）。"""
    store = get_run_store()
    run = store.get_run(run_id)
    if not run:
        raise HTTPException(status_code=404, detail=f"run {run_id} 不存在")
    out = {k: run[k] for k in ("run_id", "session_id", "status", "user_message", "created_at", "updated_at")}
    pending_json = run.get("pending_json")
    out["pending"] = json.loads(pending_json) if pending_json else None
    return out
@router.get("/api/run/{run_id}/events")
async def get_run_events(run_id: str, since: int = Query(default=0, ge=0)):
    """断线补拉：返回 seq > since 的事件（升序）。内存缓冲优先，回退 DB。"""
    store = get_run_store()
    if not store.get_run(run_id):
        raise HTTPException(status_code=404, detail=f"run {run_id} 不存在")
    bus = None
    if _main.task_executor is not None and _main.task_executor.run_engine is not None:
        bus = _main.task_executor.run_engine.bus
    if bus is not None:
        events = await bus.get_history(run_id, since)
    else:
        events = store.get_events(run_id, since)
    latest_seq = max([e.get("seq", 0) for e in events] or [since])
    return {"run_id": run_id, "events": events, "latest_seq": latest_seq}
@router.post("/api/run/{run_id}/cancel")
async def cancel_run(run_id: str):
    """取消 run：置取消标记，引擎在下一个步骤中断点响应并停止（结果丢弃不写 workspace）。"""
    store = get_run_store()
    run = store.get_run(run_id)
    if not run:
        raise HTTPException(status_code=404, detail=f"run {run_id} 不存在")
    if run["status"] in ("completed", "failed", "cancelled"):
        return {"run_id": run_id, "status": run["status"], "cancelled": False,
                "message": "run 已处于终态，无需取消"}
    store.set_cancelled(run_id)
    logger.info(f"[run] cancel requested: {run_id}")
    return {"run_id": run_id, "status": "cancelled", "cancelled": True,
            "message": "取消标记已设置，run 将在当前步骤完成后停止"}
@router.get("/suggestions")
async def get_suggestions():
    return {
        "suggestions": [
            "清除所有地图标记",
            "切换到卫星图层",
            "加载潢河郝楼可采区的矢量数据",
            "统计数据库中的点数量",
            "显示所有采样点并标注高程"
        ]
    }
class SessionCreate(BaseModel):
    title: Optional[str] = None
class SessionRename(BaseModel):
    title: str
@router.get("/api/sessions")
async def list_sessions():
    try:
        with contextlib.closing(get_db()) as conn:
            rows = conn.execute(
                "SELECT id, title, created_at, updated_at FROM sessions ORDER BY updated_at DESC"
            ).fetchall()
        return [dict(r) for r in rows]
    except Exception as e:
        logger.error(f"List sessions error: {e}")
        raise HTTPException(status_code=500, detail=str(e))
@router.post("/api/sessions")
async def create_session(body: SessionCreate):
    try:
        sid = str(uuid.uuid4())
        now = now_iso()
        title = (body.title or "新对话")[:50]
        with contextlib.closing(get_db()) as conn:
            conn.execute(
                "INSERT INTO sessions (id, title, created_at, updated_at) VALUES (?,?,?,?)",
                (sid, title, now, now)
            )
            conn.commit()
        return {"id": sid, "title": title, "created_at": now, "updated_at": now}
    except Exception as e:
        logger.error(f"Create session error: {e}")
        raise HTTPException(status_code=500, detail=str(e))
@router.delete("/api/sessions/{session_id}")
async def delete_session(session_id: str):
    try:
        with contextlib.closing(get_db()) as conn:
            conn.execute("PRAGMA foreign_keys=ON")
            conn.execute("DELETE FROM sessions WHERE id=?", (session_id,))
            conn.commit()
        return {"success": True}
    except Exception as e:
        logger.error(f"Delete session error: {e}")
        raise HTTPException(status_code=500, detail=str(e))
@router.patch("/api/sessions/{session_id}")
async def rename_session(session_id: str, body: SessionRename):
    try:
        title = body.title[:50]
        with contextlib.closing(get_db()) as conn:
            conn.execute(
                "UPDATE sessions SET title=?, updated_at=? WHERE id=?",
                (title, now_iso(), session_id)
            )
            conn.commit()
        return {"success": True, "title": title}
    except Exception as e:
        logger.error(f"Rename session error: {e}")
        raise HTTPException(status_code=500, detail=str(e))
@router.get("/api/sessions/{session_id}/messages")
async def get_session_messages(session_id: str):
    try:
        with contextlib.closing(get_db()) as conn:
            rows = conn.execute(
                "SELECT role, content, created_at FROM messages WHERE session_id=? ORDER BY id ASC",
                (session_id,)
            ).fetchall()
        return [dict(r) for r in rows]
    except Exception as e:
        logger.error(f"Get session messages error: {e}")
        raise HTTPException(status_code=500, detail=str(e))
def _schedule_fact_extraction(session_id: str, user_text: str, assistant_text: str):
    """run 成功收尾后 fire-and-forget 抽取用户事实记忆（阶段D）。同步/异步上下文通吃。"""
    try:
        from agents.fact_memory import extract_and_store_async
        llm = getattr(_main.task_executor, "_llm", None) if _main.task_executor is not None else None
        extract_and_store_async(session_id, user_text, assistant_text, llm)
    except Exception as e:  # noqa: BLE001
        logger.warning(f"[fact-memory] 调度失败: {e}")
def _persist_chat(session_id: str, user_content: str, assistant_content: str):
    """将本轮对话持久化到 SQLite，如果 session_id 不存在则跳过。"""
    if not session_id or session_id == "default":
        return
    try:
        with contextlib.closing(get_db()) as conn:
            row = conn.execute("SELECT id FROM sessions WHERE id=?", (session_id,)).fetchone()
            if not row:
                return  # 不是有效会话，跳过
            now = now_iso()
            conn.execute(
                "INSERT INTO messages (session_id, role, content, created_at) VALUES (?,?,?,?)",
                (session_id, "user", user_content, now)
            )
            conn.execute(
                "INSERT INTO messages (session_id, role, content, created_at) VALUES (?,?,?,?)",
                (session_id, "assistant", assistant_content, now)
            )
            # 自动命名：如果是第一条消息，用前20字作为标题
            msg_count = conn.execute(
                "SELECT COUNT(*) FROM messages WHERE session_id=?", (session_id,)
            ).fetchone()[0]
            if msg_count <= 2:  # 刚插入的这两条
                title = user_content[:20] + ("…" if len(user_content) > 20 else "")
                conn.execute(
                    "UPDATE sessions SET title=?, updated_at=? WHERE id=?",
                    (title, now, session_id)
                )
            else:
                conn.execute(
                    "UPDATE sessions SET updated_at=? WHERE id=?",
                    (now, session_id)
                )
            conn.commit()
    except Exception as e:
        logger.warning(f"_persist_chat failed (non-fatal): {e}")
def clean_response_content(content: str) -> str:
    if not content:
        return content
    content = re.sub(r'<\w+_tool[^>]*>', '', content)
    content = re.sub(r'\n\s*\n+', '\n\n', content)
    return content.strip()
def _optimize_map_commands(user_content: str, map_commands: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    if not map_commands:
        return []

    if not _is_explicit_map_request(user_content):
        logger.info("Suppressing map commands for non-map visualization request.")
        return []

    explicit_marker = _is_explicit_marker_request(user_content)
    seen = set()
    unique_commands = []
    for cmd in map_commands:
        key = _stable_json_key(cmd)
        if key in seen:
            continue
        seen.add(key)
        unique_commands.append(cmd)

    switch_layer_cmd = None
    set_view_cmd = None
    vector_commands = []
    marker_commands = []
    fit_markers_cmd = None
    other_commands = []

    for cmd in unique_commands:
        cmd_type = cmd.get("type")
        if cmd_type == "switch_layer":
            switch_layer_cmd = cmd
        elif cmd_type == "set_view" and set_view_cmd is None:
            set_view_cmd = cmd
        elif cmd_type == "load_vector_layer":
            vector_commands.append(cmd)
        elif cmd_type == "add_marker":
            if explicit_marker and len(marker_commands) < 3:
                marker_commands.append(cmd)
        elif cmd_type == "fit_markers":
            fit_markers_cmd = cmd
        else:
            other_commands.append(cmd)

    optimized = []
    if switch_layer_cmd:
        optimized.append(switch_layer_cmd)

    if vector_commands:
        optimized.extend(vector_commands[:2])
    else:
        if set_view_cmd:
            optimized.append(set_view_cmd)
        optimized.extend(marker_commands)
        if fit_markers_cmd and marker_commands:
            optimized.append(fit_markers_cmd)

    optimized.extend(other_commands)

    logger.info(
        "Optimized map commands from %s to %s (explicit_marker=%s)",
        len(map_commands), len(optimized), explicit_marker
    )
    return optimized
def _optimize_charts(charts: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    if not charts:
        return []

    unique = []
    seen = set()
    for chart in charts:
        config = chart.get("config") or {}
        title = ((config.get("title") or {}).get("text") if isinstance(config.get("title"), dict) else None) or chart.get("summary") or chart.get("chart_type")
        key = f"{chart.get('chart_type', 'chart')}::{title}::{_stable_json_key(config)}"
        if key in seen:
            continue
        seen.add(key)
        unique.append(chart)

    if len(unique) <= 2:
        return unique

    selected = []
    used_types = set()
    for chart in unique:
        chart_type = chart.get("chart_type", "chart")
        if chart_type not in used_types:
            selected.append(chart)
            used_types.add(chart_type)
        if len(selected) >= 2:
            break

    if not selected:
        selected = unique[:2]

    logger.info("Optimized charts from %s to %s", len(charts), len(selected))
    return selected
